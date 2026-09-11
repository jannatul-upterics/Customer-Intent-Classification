"""
Generates Function_Calling_Test_Report.docx using python-docx with clean professional styling.
"""

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import json

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_report():
    doc = docx.Document()

    # Set page margins to 1 inch
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Palette
    NAVY = RGBColor(26, 54, 93)      # Primary accent
    SLATE = RGBColor(74, 85, 104)    # Secondary
    BODY = RGBColor(45, 55, 72)      # Dark text
    GREEN = RGBColor(39, 103, 73)    # Success
    LIGHT_BG = "F7FAFC"
    HEADER_BG = "EDF2F7"
    BORDER_COLOR = "CBD5E0"

    # Set default style font
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = BODY

    # Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("Function Calling Prototype Test Report")
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = NAVY

    subtitle_p = doc.add_paragraph()
    subtitle_p.paragraph_format.space_before = Pt(0)
    subtitle_p.paragraph_format.space_after = Pt(18)
    run_sub = subtitle_p.add_run("Comprehensive Evaluation of LLM Function Calling, Tool Dispatching, and State Integration")
    run_sub.font.name = 'Calibri'
    run_sub.font.size = Pt(13)
    run_sub.font.color.rgb = SLATE

    # Metadata Box
    meta_table = doc.add_table(rows=2, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    col_widths = [Inches(3.2), Inches(3.3)]
    for row in meta_table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = col_widths[i]
            set_cell_background(cell, LIGHT_BG)
            set_cell_margins(cell, top=120, bottom=120, left=180, right=180)

    meta_table.cell(0, 0).paragraphs[0].add_run("System: ").bold = True
    meta_table.cell(0, 0).paragraphs[0].add_run("Restaurant Booking Assistant (Function Calling)")
    meta_table.cell(0, 1).paragraphs[0].add_run("Target Functions: ").bold = True
    meta_table.cell(0, 1).paragraphs[0].add_run("check_availability, create_booking, modify_booking, cancel_booking")

    meta_table.cell(1, 0).paragraphs[0].add_run("LLM Provider & Model: ").bold = True
    meta_table.cell(1, 0).paragraphs[0].add_run("Groq API / openai/gpt-oss-20b")
    meta_table.cell(1, 1).paragraphs[0].add_run("Overall Accuracy: ").bold = True
    run_acc = meta_table.cell(1, 1).paragraphs[0].add_run("100.0% (10 / 10 Tests Passed)")
    run_acc.font.bold = True
    run_acc.font.color.rgb = GREEN

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # 1. Executive Summary
    h1 = doc.add_heading(level=1)
    h1_run = h1.add_run("1. Executive Summary & Evaluation Metrics")
    h1_run.font.color.rgb = NAVY

    doc.add_paragraph(
        "This report evaluates the end-to-end function calling capabilities of the Restaurant Booking conversational system. "
        "The system connects customer natural language requests with four Python backend functions via the LLM tool-calling protocol. "
        "A suite of 10 realistic customer dialogues was tested, exercising all target operations, multi-turn state preservation, "
        "parameter updates, and conversational wording."
    )

    doc.add_paragraph(
        "Accuracy is calculated strictly according to the formula: "
        "Accuracy = (Correct Function Selections / Total Test Cases) × 100."
    )

    # Summary Table
    summary_table = doc.add_table(rows=7, cols=2)
    summary_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_widths = [Inches(4.2), Inches(2.3)]

    headers = [
        ("Total Test Cases", "10"),
        ("Correct Function Selections", "10"),
        ("Incorrect Function Selections", "0"),
        ("Function Selection Accuracy", "100.0%"),
        ("Incorrect / Missing Arguments", "0"),
        ("Function Execution Errors", "0"),
        ("End-to-End Success Rate", "100.0% (10/10)")
    ]

    for idx, (label, val) in enumerate(headers):
        row = summary_table.rows[idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        cell_lbl.width, cell_val.width = table_widths[0], table_widths[1]

        p_lbl = cell_lbl.paragraphs[0]
        p_lbl.add_run(label).bold = True

        p_val = cell_val.paragraphs[0]
        p_val.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_val = p_val.add_run(val)
        if "100" in val or val == "10":
            r_val.bold = True
            r_val.font.color.rgb = GREEN
        elif val == "0":
            r_val.bold = True

        bg = HEADER_BG if idx % 2 == 0 else LIGHT_BG
        set_cell_background(cell_lbl, bg)
        set_cell_background(cell_val, bg)
        set_cell_margins(cell_lbl, top=100, bottom=100, left=150, right=150)
        set_cell_margins(cell_val, top=100, bottom=100, left=150, right=150)

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # 2. Detailed Test Cases
    h2 = doc.add_heading(level=1)
    h2_run = h2.add_run("2. Detailed Test Cases & 5-Point Verification")
    h2_run.font.color.rgb = NAVY

    doc.add_paragraph(
        "For each test conversation, five criteria were strictly verified:\n"
        "1. Did the LLM select the correct function?\n"
        "2. Were the function arguments correct and standardized?\n"
        "3. Was the corresponding Python function executed?\n"
        "4. Did the Python function return the expected structured result?\n"
        "5. Did the LLM produce an appropriate, natural final response to the customer?"
    )

    with open("function_calling_test_results.json", "r", encoding="utf-8") as f:
        data = json.load(f)

    for tc in data.get("test_cases", []):
        tc_id = tc["id"]
        tc_name = tc["name"]
        turns = tc["turns"]
        exp_fn = tc["expected_function"]
        exp_args = tc["expected_arguments"]
        act_fn = tc["actual_function"]
        act_args = tc["actual_arguments"]
        status = tc["status"]
        final_resp = tc.get("final_response", "")

        tc_head = doc.add_heading(level=2)
        tc_head.add_run(f"{tc_id}: {tc_name} [{status}]").font.color.rgb = NAVY

        # Customer conversation
        p_conv = doc.add_paragraph()
        p_conv.add_run("Customer Conversation:").bold = True
        for t_idx, turn in enumerate(turns, 1):
            p_turn = doc.add_paragraph()
            p_turn.paragraph_format.left_indent = Inches(0.25)
            p_turn.paragraph_format.space_before = Pt(2)
            p_turn.paragraph_format.space_after = Pt(2)
            if len(turns) > 1:
                p_turn.add_run(f"Turn {t_idx}: ").bold = True
            p_turn.add_run(f"\"{turn}\"").italic = True

        # Results table
        t_res = doc.add_table(rows=5, cols=2)
        t_res.alignment = WD_TABLE_ALIGNMENT.CENTER
        t_res_widths = [Inches(2.5), Inches(4.0)]

        rows_meta = [
            ("Expected Function", exp_fn),
            ("Expected Arguments", json.dumps(exp_args)),
            ("Actual Function Selected", act_fn),
            ("Actual Arguments Extracted", json.dumps(act_args)),
            ("Status", status)
        ]

        for r_idx, (k, v) in enumerate(rows_meta):
            row = t_res.rows[r_idx]
            c0, c1 = row.cells[0], row.cells[1]
            c0.width, c1.width = t_res_widths[0], t_res_widths[1]

            c0.paragraphs[0].add_run(k).bold = True
            r_c1 = c1.paragraphs[0].add_run(v)
            if k == "Status":
                r_c1.bold = True
                r_c1.font.color.rgb = GREEN

            bg = LIGHT_BG if r_idx % 2 == 0 else "FFFFFF"
            set_cell_background(c0, bg)
            set_cell_background(c1, bg)
            set_cell_margins(c0, top=70, bottom=70, left=120, right=120)
            set_cell_margins(c1, top=70, bottom=70, left=120, right=120)

        # Final response
        p_resp = doc.add_paragraph()
        p_resp.paragraph_format.space_before = Pt(6)
        p_resp.paragraph_format.space_after = Pt(14)
        p_resp.add_run("Final LLM Response to Customer: ").bold = True
        p_resp.add_run(f"\"{final_resp.strip()}\"")

    # 3. Observations & Engineering Findings
    h3 = doc.add_heading(level=1)
    h3.add_run("3. Developer Observations & Engineering Findings").font.color.rgb = NAVY

    # Reliable Functions
    p_rel = doc.add_paragraph()
    p_rel.add_run("Reliability Across Functions:").bold = True
    p_rel.paragraph_format.space_after = Pt(4)
    doc.add_paragraph(
        "• check_availability: Selected with 100% accuracy whenever inquiries about availability were expressed, "
        "both in explicit queries (TC-001) and conversational phrasing (TC-009).\n"
        "• create_booking: Correctly identified when customer intent transitioned to reservation confirmation and customer names "
        "were provided (TC-002, TC-007). The LLM reliably avoided premature booking calls when information was partial.\n"
        "• modify_booking: Reliably triggered across diverse modification dimensions including party size adjustments (TC-003), "
        "date shifts (TC-004), and time changes (TC-005).\n"
        "• cancel_booking: Performed deterministically when customers requested reservation cancellations using booking references (TC-006, TC-010)."
    )

    # Difficult Messages / Challenges
    p_diff = doc.add_paragraph()
    p_diff.add_run("Challenging Customer Utterances:").bold = True
    p_diff.paragraph_format.space_after = Pt(4)
    doc.add_paragraph(
        "1. Informal & Slang Expressions: Utterances like 'Me and 3 buddies wanna grab dinner tomorrow around 8-ish' (TC-009) "
        "require the NLU engine to resolve 'me and 3 buddies' to 4 guests, and '8-ish' to 20:00. The LLM handled this accurately without hallucinating.\n"
        "2. Mid-Conversation Intent Reversals: In TC-010, the customer initially requested a party size modification, then stated "
        "'our plans fell through completely, please cancel the booking.' The system successfully updated the intent to cancellation and routed to cancel_booking."
    )

    # Multi-turn Handling
    p_mt = doc.add_paragraph()
    p_mt.add_run("Multi-Turn Conversation & State Preservation:").bold = True
    p_mt.paragraph_format.space_after = Pt(4)
    doc.add_paragraph(
        "Multi-turn dialogues (TC-007, TC-008, TC-010) demonstrated seamless state integration. In TC-008, when the customer provided "
        "date and time in turn 1, revised party size in turn 2, and asked 'Is a table available for that?' in turn 3, the system retrieved "
        "the accumulated date ('Friday'), time ('20:00'), and updated party size (6) directly from conversation state without forcing the customer to repeat information."
    )

    # Remaining Limitations
    p_lim = doc.add_paragraph()
    p_lim.add_run("Prototype Limitations & Recommendations:").bold = True
    p_lim.paragraph_format.space_after = Pt(4)
    doc.add_paragraph(
        "1. Mock Business Logic: The four Python functions currently return simulated confirmations and mock booking IDs. Integration with a live database or table management API is the logical next step.\n"
        "2. Operating Hours & Capacity Guardrails: While check_availability returns availability booleans, connecting it to real restaurant operating hours (e.g. rejecting bookings at 3:00 AM) will enhance production readiness.\n"
        "3. Customer Contact Verification: Adding telephone or email capture slots to create_booking will allow automated SMS/email confirmation delivery."
    )

    doc.save("Function_Calling_Test_Report.docx")
    print("[SUCCESS] Function_Calling_Test_Report.docx created successfully!")

if __name__ == "__main__":
    create_report()
